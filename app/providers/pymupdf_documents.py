"""PyMuPDF and python-pptx document extraction provider for VirtuJudge AI-ML.

Extracts text from PDF and PPTX documents with page/slide provenance, splits into
versioned chunks, and optionally generates text embeddings via an EmbeddingProvider.
Treats all document text strictly as untrusted data.
"""

import asyncio
import hashlib
import logging
from pathlib import Path

import pymupdf
from pptx import Presentation

from app.providers.base import EmbeddingProvider
from app.providers.types import DocumentChunk

logger = logging.getLogger(__name__)

DEFAULT_CHUNK_SIZE = 512
DEFAULT_CHUNK_OVERLAP = 64
CHUNKING_VERSION = "v1_512_64"


class DocumentExtractionError(Exception):
    """Exception raised when document extraction fails."""


class PyMuPDFDocumentProvider:
    """Document extraction provider powered by PyMuPDF (PDF) and python-pptx (PPTX)."""

    def __init__(
        self,
        embedding_provider: EmbeddingProvider | None = None,
        *,
        chunk_size: int = DEFAULT_CHUNK_SIZE,
        chunk_overlap: int = DEFAULT_CHUNK_OVERLAP,
        chunking_version: str = CHUNKING_VERSION,
    ) -> None:
        self.embedding_provider = embedding_provider
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.chunking_version = chunking_version

    def _chunk_text(
        self,
        text: str,
        asset_version_id: str,
        page_or_slide: int,
        extraction_method: str,
    ) -> list[DocumentChunk]:
        """Split page/slide text into versioned, deterministic chunks."""
        text = text.strip()
        if not text:
            return []

        chunks: list[DocumentChunk] = []
        text_len = len(text)

        # The stepping loop naturally handles single-chunk texts (when text_len <= chunk_size)
        # as well as multi-chunk sliding windows with identical offset and hash semantics.
        step = max(1, self.chunk_size - self.chunk_overlap)
        start_offset = 0

        while start_offset < text_len:
            end_offset = min(text_len, start_offset + self.chunk_size)
            chunk_content = text[start_offset:end_offset]

            chunk_key = f"{asset_version_id}:{page_or_slide}:{start_offset}:{self.chunking_version}"
            chunk_hash = hashlib.sha256(chunk_key.encode()).hexdigest()[:16]
            chunk_id = f"{asset_version_id}_p{page_or_slide}_{chunk_hash}"

            chunks.append(
                DocumentChunk(
                    chunk_id=chunk_id,
                    asset_version_id=asset_version_id,
                    page_or_slide=page_or_slide,
                    text=chunk_content,
                    start_offset=start_offset,
                    end_offset=end_offset,
                    extraction_method=extraction_method,
                    chunking_version=self.chunking_version,
                )
            )

            if end_offset >= text_len:
                break
            start_offset += step

        return chunks

    def _extract_pdf(self, doc_path: Path, asset_version_id: str) -> list[DocumentChunk]:
        """Extract text and chunk by page from a PDF file."""
        extraction_method = f"pymupdf/{pymupdf.__version__}"
        try:
            doc = pymupdf.open(str(doc_path))  # type: ignore[no-untyped-call]
        except Exception as exc:
            raise DocumentExtractionError(f"Failed to open PDF document: {exc}") from exc

        try:
            if doc.is_encrypted and not doc.authenticate(""):  # type: ignore[no-untyped-call]
                raise DocumentExtractionError(
                    "PDF document is password-protected or encrypted."
                )

            all_chunks: list[DocumentChunk] = []
            for page_index in range(len(doc)):
                page = doc[page_index]
                page_text = page.get_text() or ""  # type: ignore[no-untyped-call]
                page_chunks = self._chunk_text(
                    page_text,
                    asset_version_id=asset_version_id,
                    page_or_slide=page_index + 1,
                    extraction_method=extraction_method,
                )
                all_chunks.extend(page_chunks)
            return all_chunks
        finally:
            doc.close()  # type: ignore[no-untyped-call]

    def _extract_pptx(self, doc_path: Path, asset_version_id: str) -> list[DocumentChunk]:
        """Extract text and chunk by slide from a PPTX file."""
        extraction_method = "python-pptx"
        try:
            prs = Presentation(str(doc_path))
        except Exception as exc:
            raise DocumentExtractionError(f"Failed to open PPTX presentation: {exc}") from exc

        all_chunks: list[DocumentChunk] = []
        for slide_index, slide in enumerate(prs.slides):
            slide_texts: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = "".join(run.text for run in paragraph.runs).strip()
                        if text:
                            slide_texts.append(text)
                elif shape.has_table:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            c_text = cell.text.strip()
                            if c_text:
                                slide_texts.append(c_text)

            slide_text = "\n".join(slide_texts).strip()
            slide_chunks = self._chunk_text(
                slide_text,
                asset_version_id=asset_version_id,
                page_or_slide=slide_index + 1,
                extraction_method=extraction_method,
            )
            all_chunks.extend(slide_chunks)
        return all_chunks

    def _sync_extract(self, doc_path: Path, asset_version_id: str) -> list[DocumentChunk]:
        """Synchronously parse and extract document chunks based on extension."""
        if not doc_path.is_file():
            raise FileNotFoundError(f"Document file not found: {doc_path}")

        suffix = doc_path.suffix.lower()
        if suffix == ".pdf":
            return self._extract_pdf(doc_path, asset_version_id)
        elif suffix == ".pptx":
            return self._extract_pptx(doc_path, asset_version_id)
        else:
            raise DocumentExtractionError(
                f"Unsupported document format '{suffix}'. Only .pdf and .pptx are supported."
            )

    async def extract_and_embed(
        self, doc_path: Path, *, asset_version_id: str = ""
    ) -> list[DocumentChunk]:
        """Extract chunks from PDF/PPTX in thread pool and embed if provider is set."""
        resolved_version = asset_version_id or doc_path.stem
        chunks = await asyncio.to_thread(self._sync_extract, Path(doc_path), resolved_version)

        if not chunks:
            return []

        if self.embedding_provider is not None:
            texts = [chunk.text for chunk in chunks]
            try:
                embeddings = await self.embedding_provider.embed(texts)
                model_name = getattr(self.embedding_provider, "model_name", None)
                for chunk, emb in zip(chunks, embeddings, strict=True):
                    chunk.embedding = emb
                    chunk.embedding_dimensions = len(emb)
                    chunk.embedding_model = model_name
            except Exception as exc:
                # Log safe identifier only, never the host file path or document content
                logger.warning(
                    "Embedding generation failed for asset_version_id=%s: %s",
                    resolved_version,
                    exc,
                )
                raise DocumentExtractionError(f"Embedding generation failed: {exc}") from exc

        return chunks


__all__ = [
    "CHUNKING_VERSION",
    "DEFAULT_CHUNK_OVERLAP",
    "DEFAULT_CHUNK_SIZE",
    "DocumentExtractionError",
    "PyMuPDFDocumentProvider",
]
